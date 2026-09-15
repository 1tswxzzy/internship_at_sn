import os
import sys
import re
import numpy as np
from collections import OrderedDict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==================== 工具函数 ====================

def append_district_sheets_to_file(result_file_path, stats_rows, summary_rows):
    """将区县差异系数和汇总 sheet 追加到已有的分析结果文件"""
    try:
        wb = openpyxl.load_workbook(result_file_path)
    except Exception as e:
        print(f"无法打开结果文件以追加区县统计：{e}")
        return

    # 如果已有同名 sheet，先删除
    for name in ['区县差异系数统计', '区县汇总']:
        if name in wb.sheetnames:
            del wb[name]

    if stats_rows:
        ws_stats = wb.create_sheet('区县差异系数统计')
        for row in stats_rows:
            ws_stats.append(row)

    if summary_rows:
        ws_summary = wb.create_sheet('区县汇总')
        for row in summary_rows:
            ws_summary.append(row)

    wb.save(result_file_path)
    print(f"区县差异系数已添加至：{result_file_path}")

def analyze_district_from_result_file(result_file_path, school_type):
    """
    从已生成的分析结果文件中读取各指标 sheet，按区县分组，计算差异系数。
    注意：结果表中的指标值已经是每百名或生均值，因此直接以学生数加权计算。
    返回: stats_rows, summary_rows (两个二维列表，用于写入区县差异系数统计和区县汇总)
    """
    import openpyxl
    from collections import OrderedDict
    import numpy as np

    wb = openpyxl.load_workbook(result_file_path, data_only=True)
    configs = [
        ('每百名学生多媒体教室数', 'ind1', True),
        ('每百名学生高学历教师数', 'ind2', False),
        ('每百名学生骨干教师数', 'ind3', False),
        ('每百名学生艺体教师数', 'ind4', False),
        ('生均教学及辅助用房面积', 'ind5', True),
        ('生均运动场馆面积', 'ind6', True),
        ('生均教学仪器设备值', 'ind7', True),
    ]

    school_values = OrderedDict()
    for sheet_name, ind_key, use_adjusted in configs:
        if sheet_name not in wb.sheetnames:
            print(f"警告：结果文件中没有 sheet '{sheet_name}'，跳过。")
            continue
        ws = wb[sheet_name]
        headers = [cell.value for cell in ws[1]]  # 第一行是标题

        # 查找必需的列索引
        try:
            school_col = headers.index('学校名称') + 1
            district_col = headers.index('统计三级') + 1   # 或 '所在地区划三级'，视实际表头而定
            # 查找学生数列（可能名称略有不同，简单匹配“在校生数”）
            stu_col = next(i+1 for i, h in enumerate(headers) if h and '在校生数' in str(h))
        except (ValueError, StopIteration):
            print(f"Sheet '{sheet_name}' 缺少必要列（学校名称、统计三级、在校生数），跳过。")
            continue

        # 查找目标指标列
        if use_adjusted:
            target_col_name = '调整后指标值'
        else:
            if '高学历' in sheet_name:
                target_col_name = '每百名学生高学历教师数'
            elif '骨干' in sheet_name:
                target_col_name = '每百名学生骨干教师数'
            elif '艺体' in sheet_name:
                target_col_name = '每百名学生艺体教师数'
            else:
                continue
        if target_col_name not in headers:
            print(f"Sheet '{sheet_name}' 中缺少列 '{target_col_name}'，跳过。")
            continue
        target_col = headers.index(target_col_name) + 1

        # 读取数据行
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(row) < max(school_col, district_col, stu_col, target_col):
                continue
            school = str(row[school_col-1]).strip() if row[school_col-1] else ''
            district = str(row[district_col-1]).strip() if row[district_col-1] else ''
            if not school:
                continue
            try:
                stu = float(row[stu_col-1]) if row[stu_col-1] is not None else 0
            except (ValueError, TypeError):
                stu = 0
            try:
                ind_val = float(row[target_col-1]) if row[target_col-1] is not None else 0
            except (ValueError, TypeError):
                ind_val = 0

            if school not in school_values:
                school_values[school] = {'district': district, 'A': stu, 'inds': [0]*7}
            # 学生数以第一次读取的为准（通常各 sheet 的学生数一致）
            if school_values[school]['A'] == 0:
                school_values[school]['A'] = stu
            idx = int(ind_key[-1]) - 1
            school_values[school]['inds'][idx] = ind_val

    # 按区县分组
    district_data = OrderedDict()
    for school, data in school_values.items():
        district = data['district']
        if not district:
            district = '未知区县'
        if district not in district_data:
            district_data[district] = []
        district_data[district].append({
            'school': school,
            'A': data['A'],
            'ind1': data['inds'][0],
            'ind2': data['inds'][1],
            'ind3': data['inds'][2],
            'ind4': data['inds'][3],
            'ind5': data['inds'][4],
            'ind6': data['inds'][5],
            'ind7': data['inds'][6],
        })

    if not district_data:
        print("未提取到任何区县数据，无法生成差异系数。")
        return None, None

    # ==================== 计算统计量 ====================
    # 指标名称列表（不再需要乘数）
    indicator_names = [
        '多媒体教室',
        '高学历教师',
        '骨干教师',
        '艺体教师',
        '教学及辅助用房净面积',
        '运动场馆面积',
        '教学仪器设备值'
    ]

    def calc_stats(data_list, ind_key, district_name, indicator_name):
        """计算加权平均值、标准差、差异系数（直接使用结果表中的指标值）"""
        print(f"\n{'='*60}")
        print(f"📊 计算区县【{district_name}】-【{indicator_name}】")
        print(f"{'='*60}")
        print(f"{'学校名称':<30} {'学生数(A)':>10} {'指标值':>12}")
        print("-" * 60)

        A_sum = 0.0
        weighted_sum = 0.0   # Σ(指标值 × 学生数)
        values = []           # 存储 (学生数, 指标值)

        for d in data_list:
            A = d['A']
            ind = d[ind_key]
            A_sum += A
            weighted_sum += ind * A
            values.append((A, ind))
            print(f"{d['school']:<30} {A:>10.1f} {ind:>12.4f}")

        print("-" * 60)
        print(f"  学生总数(A_sum) = {A_sum:.2f}")
        print(f"  加权总和 = {weighted_sum:.2f}")

        if A_sum == 0:
            print("  → 学生总数为0，无法计算，返回 X=0, S=0, CV=0")
            return {'X': 0, 'S': 0, 'CV': 0}

        # 加权平均值
        X = weighted_sum / A_sum

        # 加权方差
        variance = 0.0
        for A, ind in values:
            variance += (A / A_sum) * (ind - X) ** 2

        S = np.sqrt(variance) if variance > 0 else 0.0
        CV = round((S / X), 2) if X != 0 else 0.0

        print(f"  ✨ 计算结果：")
        print(f"     加权平均值 X = {weighted_sum:.2f} / {A_sum:.2f} = {X:.4f}")
        print(f"     加权方差   = {variance:.6f}")
        print(f"     加权标准差 S = sqrt(方差) = {S:.6f}")
        print(f"     差异系数 CV = {CV:.2f}")
        print(f"{'='*60}\n")

        return {'X': round(X, 4), 'S': round(S, 6), 'CV': CV}

    # 分组（分配组别 T1, T2...）
    group_counter = 1
    district_groups = OrderedDict()
    for dist in district_data:
        district_groups[dist] = {
            'group': f"T{group_counter}",
            'data': district_data[dist]
        }
        group_counter += 1

    # 构建差异系数表头
    header_parts = ['组别', '区县']
    for name in indicator_names:
        header_parts.append(f'{name}平均值')
        header_parts.append(f'{name}标准差')
        header_parts.append(f'{name}差异系数')
    stats_rows = [header_parts]

    # 汇总各指标的所有数据（用于总计）
    all_data_lists = {f'ind{i+1}': [] for i in range(7)}
    for dist, info in district_groups.items():
        data = info['data']
        row = [info['group'], dist]
        for i, name in enumerate(indicator_names):
            ind_key = f'ind{i+1}'
            stats = calc_stats(data, ind_key, dist, name)   # 不再传乘数
            row.extend([stats['X'], stats['S'], stats['CV']])
        stats_rows.append(row)
        for i in range(7):
            all_data_lists[f'ind{i+1}'].extend(data)

    # 总计行
    print("\n" + "=" * 60)
    print("             🌟 计算总体（总计）汇总统计 🌟")
    print("=" * 60)
    total_row = ['', '总计']
    for i, name in enumerate(indicator_names):
        stats = calc_stats(all_data_lists[f'ind{i+1}'], f'ind{i+1}', '总计', name)
        total_row.extend([stats['X'], stats['S'], stats['CV']])
    stats_rows.append(total_row)

    # ==================== 区县汇总表 ====================
    summary_header = ['区县', '学生总数(A)'] + indicator_names
    summary_rows = [summary_header]
    total_A = 0
    total_inds = [0]*7
    for dist, info in district_groups.items():
        data = info['data']
        sA = sum(d['A'] for d in data)
        row = [dist, sA]
        for i in range(7):
            si = sum(d[f'ind{i+1}'] for d in data)
            row.append(round(si, 2))
            total_inds[i] += si
        summary_rows.append(row)
        total_A += sA
    summary_rows.append(['总计', total_A] + [round(v, 2) for v in total_inds])

    wb.close()
    return stats_rows, summary_rows

def classify_school_type(file_path):
    """根据文件名判断学校类型（小学/中学/初中/高中）"""
    file_name = os.path.basename(file_path)
    if '小学' in file_name:
        return '小学'
    elif '初中' in file_name:
        return '初中'
    elif '高中' in file_name:
        return '高中'
    elif '中学' in file_name:
        return '中学'
    return '未知'

def find_district_col(ws, head_size):
    """查找“统计三级”列"""
    col = find_col_by_keywords(ws, head_size, '统计三级')
    if not col:
        col = find_col_by_hierarchy(ws, head_size, ['统计三级'])
    return col

def extract_district(text):
    if not text or str(text).strip() == '':
        return None
    text = str(text).strip()
    pattern = r'([^\s,，、]+(?:区|县|市|自治县|自治旗))'
    match = re.search(pattern, text)
    if match:
        return match.group(1)
    return None

def find_col_by_keywords(ws, head_size, keywords, match_mode='any'):
    """通过关键词匹配列，优先底层表头"""
    if isinstance(keywords, str):
        keywords = [keywords]
    found_cols = set()
    for row_idx in range(head_size, 0, -1):
        for col_idx in range(1, ws.max_column + 1):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val and isinstance(cell_val, str):
                for kw in keywords:
                    if kw in cell_val:
                        if match_mode == 'any':
                            return col_idx
                        found_cols.add(col_idx)
        if match_mode == 'sum' and len(found_cols) >= len(keywords):
            return list(found_cols)
        if match_mode == 'all' and len(found_cols) == len(keywords):
            break
    if match_mode == 'any' and not found_cols:
        return None
    if match_mode == 'all':
        return found_cols.pop() if found_cols else None
    if match_mode == 'sum':
        return list(found_cols)
    return None

def find_col_by_hierarchy(ws, head_size, hierarchy):
    """根据表头层级路径查找唯一列"""
    for col in range(1, ws.max_column + 1):
        idx = 0
        for row in range(1, head_size + 1):
            cell_val = ws.cell(row=row, column=col).value
            if cell_val and isinstance(cell_val, str) and hierarchy[idx] in cell_val:
                idx += 1
                if idx == len(hierarchy):
                    return col
    return None

def find_cols_with_prefix(ws, head_size, prefix_hierarchy):
    """返回所有满足前缀层级序列的列号"""
    cols = []
    for col in range(1, ws.max_column + 1):
        idx = 0
        for row in range(1, head_size + 1):
            cell_val = ws.cell(row=row, column=col).value
            if cell_val and isinstance(cell_val, str) and prefix_hierarchy[idx] in cell_val:
                idx += 1
                if idx == len(prefix_hierarchy):
                    cols.append(col)
                    break
    return cols

# ==================== Excel 基础读取 ====================
def get_worksheet_info(file_path):
    """打开 Excel，返回工作簿、工作表、有效行数、有效列数"""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
    except FileNotFoundError:
        print(f"错误：找不到文件 '{file_path}'")
        return None, None, None, None
    except Exception as e:
        print(f"读取文件时发生错误：{e}")
        return None, None, None, None

    max_h = ws.max_row
    while max_h > 0:
        if any(cell.value is not None for cell in ws[max_h]):
            break
        max_h -= 1

    max_l = ws.max_column
    while max_l > 0:
        if any(ws.cell(row=r, column=max_l).value is not None for r in range(1, max_h+1)):
            break
        max_l -= 1

    return wb, ws, max_h, max_l

def guess_head_size(ws, max_rows_to_check=20):
    """自动推断表头行数"""
    head_size = 1
    if ws.merged_cells.ranges:
        head_end = 0
        for mr in ws.merged_cells.ranges:
            if mr.max_row > head_end:
                head_end = mr.max_row
        if head_end > 0:
            head_end = min(head_end, ws.max_row)
            preview = [[cell.value if cell.value is not None else "" for cell in ws[row]]
                       for row in range(1, head_end + 1)]
            return head_end, preview

    data_start = 1
    for row_idx in range(1, min(max_rows_to_check, ws.max_row) + 1):
        for cell in ws[row_idx]:
            if cell.value is None:
                continue
            if isinstance(cell.value, (int, float)):
                data_start = row_idx
                break
            if isinstance(cell.value, str):
                try:
                    float(cell.value.replace(',', '').replace('%', '').strip())
                    data_start = row_idx
                    break
                except ValueError:
                    continue
        if data_start > 1:
            break

    head_size = data_start - 1 if data_start > 1 else 1
    preview = [[cell.value if cell.value is not None else "" for cell in ws[row]]
               for row in range(1, head_size + 1)]
    return head_size, preview

# ==================== 查找"是否计校数"列 ====================
def find_count_school_col(ws, head_size):
    return find_col_by_keywords(ws, head_size, '是否计校数')

def build_school_merge_map(ws, head_size):
    """构建不计校数学校的合并映射"""
    count_col = find_count_school_col(ws, head_size)
    if not count_col:
        print("未找到'是否计校数'列，跳过不计校数合并。")
        return {}

    pending_schools = set()
    for r in range(head_size + 1, ws.max_row + 1):
        try:
            if float(ws.cell(row=r, column=count_col).value) == 0:
                name = ws.cell(row=r, column=1).value
                if name:
                    pending_schools.add(str(name).strip())
        except (ValueError, TypeError):
            pass

    if not pending_schools:
        print("没有发现不计校数的学校。")
        return {}

    print(f"\n发现 {len(pending_schools)} 个不计校数的学校，请依次输入它们所属的原校点名称：")
    merge_map = {}
    for school in pending_schools:
        while True:
            master = input(f"  {school}  → 原校点名称：").strip()
            if master:
                merge_map[school] = master
                break
            print("  输入不能为空，请重新输入。")
    print("映射关系已记录。")
    return merge_map

# ==================== 特殊学校查找与学段分析 ====================
def find_special_school_types(ws, head_size, max_h, file_type):
    """查找特殊办学类型学校"""
    school_type_col = None
    for col in range(1, ws.max_column + 1):
        for row in range(1, head_size + 1):
            cell_val = ws.cell(row=row, column=col).value
            if cell_val and isinstance(cell_val, str) and '办学类型' in cell_val:
                school_type_col = col
                break
        if school_type_col:
            break
    
    if not school_type_col:
        school_type_col = find_col_by_keywords(ws, head_size, '办学类型')
    
    if not school_type_col:
        return None
    
    special_types = {
        '九年一贯制': [],
        '十二年一贯制': [],
        '完全中学': []
    }
    
    for r in range(head_size + 1, max_h + 1):
        type_val = ws.cell(row=r, column=school_type_col).value
        if type_val and isinstance(type_val, str):
            type_str = str(type_val).strip()
            for special_type in special_types.keys():
                if special_type in type_str:
                    school_name = ws.cell(row=r, column=1).value
                    if school_name:
                        school_name_str = str(school_name).strip()
                        stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
                        if not stu_col:
                            stu_col = find_col_by_keywords(ws, head_size, '在校生数')
                        stu_val = ws.cell(row=r, column=stu_col).value if stu_col else None
                        try:
                            stu_count = float(stu_val) if stu_val is not None else 0
                        except (ValueError, TypeError):
                            stu_count = 0
                        
                        special_types[special_type].append({
                            'row': r,
                            'name': school_name_str,
                            'full_type': type_str,
                            'student_count': stu_count,
                            'source_file_type': file_type
                        })
    
    total_found = 0
    for special_type, schools in special_types.items():
        if schools:
            print(f"\n{special_type}：共 {len(schools)} 所")
            for school in schools:
                print(f"  行{school['row']}: {school['name']} (在校生: {school['student_count']})")
            total_found += len(schools)
    
    return special_types

def find_school_in_file(file_path, school_name):
    """在指定文件中查找学校，返回其在校生数"""
    wb, ws, max_h, max_l = get_worksheet_info(file_path)
    if wb is None:
        return None
    
    head_size, _ = guess_head_size(ws)
    
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    
    if not stu_col:
        return None
    
    # 精确匹配
    for r in range(head_size + 1, max_h + 1):
        cell_name = ws.cell(row=r, column=1).value
        if cell_name and str(cell_name).strip() == school_name:
            stu_val = ws.cell(row=r, column=stu_col).value
            try:
                return float(stu_val) if stu_val is not None else 0
            except (ValueError, TypeError):
                return 0
    
    # 模糊匹配
    for r in range(head_size + 1, max_h + 1):
        cell_name = ws.cell(row=r, column=1).value
        if cell_name:
            name_str = str(cell_name).strip()
            if school_name in name_str or name_str in school_name:
                stu_val = ws.cell(row=r, column=stu_col).value
                try:
                    return float(stu_val) if stu_val is not None else 0
                except (ValueError, TypeError):
                    return 0
    
    return None

def calculate_section_ratios(all_special_schools, file_map):
    """
    跨文件计算特殊学校的学段占比。
    九年一贯制：小学权重1，初中权重1.1
    十二年一贯制：小学权重1，初中权重1.1，高中权重1.32
    完全中学：初中权重1，高中权重1.2
    """
    weights = {
        '九年一贯制': {'小学': 1, '初中': 1.1},
        '十二年一贯制': {'小学': 1, '初中': 1.1, '高中': 1.32},
        '完全中学': {'初中': 1, '高中': 1.2}
    }
    
    section_file_map = {'小学': '小学', '初中': '初中', '高中': '高中'}
    
    print("\n" + "=" * 60)
    print("第一步：计算特殊学校学段占比")
    print("=" * 60)
    
    section_ratios = {}
    
    for special_type, schools in all_special_schools.items():
        if not schools:
            continue
        
        school_type_weights = weights[special_type]
        print(f"\n计算【{special_type}】学校学段占比...")
        
        for school in schools:
            school_name = school['name']
            section_students = {}
            
            for section in school_type_weights.keys():
                file_type = section_file_map.get(section)
                if file_type and file_type in file_map:
                    file_path = file_map[file_type]
                    student_count = find_school_in_file(file_path, school_name)
                    if student_count is not None:
                        section_students[section] = student_count
                    else:
                        section_students[section] = 0
            
            # 计算加权值和占比
            total_weighted = 0
            weighted_values = {}
            
            for section, weight in school_type_weights.items():
                student_count = section_students.get(section, 0)
                weighted_value = student_count * weight
                total_weighted += weighted_value
                weighted_values[section] = weighted_value
            
            # 计算各学段占比
            ratios = {}
            for section in school_type_weights.keys():
                if total_weighted > 0:
                    ratios[section] = weighted_values[section] / total_weighted
                else:
                    ratios[section] = 0
            
            section_ratios[school_name] = {
                'type': special_type,
                'ratios': ratios,
                'section_students': section_students,
                'total_weighted': total_weighted
            }
            
            print(f"  {school_name}：")
            for section in school_type_weights.keys():
                print(f"    {section}：学生数={section_students.get(section, 0)}, "
                      f"权重={school_type_weights[section]}, "
                      f"加权值={weighted_values.get(section, 0):.2f}, "
                      f"占比={ratios[section]*100:.2f}%")
    
    return section_ratios

def get_section_ratio_for_school(school_name, section_ratios, target_section):
    """
    获取某所学校在目标学段的占比。
    如果学校是特殊学校，返回该学段的占比；
    如果是普通学校，返回1.0（100%）。
    """
    if school_name in section_ratios:
        ratio_info = section_ratios[school_name]
        return ratio_info['ratios'].get(target_section, 1.0)
    return 1.0

# ==================== 计算函数 ====================
# 不需要使用占比的指标（保持不变）
def calc_high_education_per_stu(ws, head_size, max_h, school_type, merge_map=None):
    """2. 每百名学生高学历教师数 - 不需要调整"""
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    if not stu_col:
        print("错误：未找到'在校生数'列。")
        return None

    edu_region = find_cols_with_prefix(ws, head_size, ['专任教师', '#按学历分'])
    if not edu_region:
        kw_primary = ['硕士研究生', '本科', '专科']
        kw_middle = ['硕士研究生', '本科']
        keywords = kw_primary if school_type == '小学' else kw_middle
        target_cols = []
        for kw in keywords:
            col = find_col_by_keywords(ws, head_size, kw)
            if col:
                target_cols.append(col)
    else:
        kw_primary = ['硕士研究生', '本科', '专科']
        kw_middle = ['硕士研究生', '本科']
        keywords = kw_primary if school_type == '小学' else kw_middle
        target_cols = []
        for kw in keywords:
            found = None
            for col in edu_region:
                for r in range(1, head_size + 1):
                    val = ws.cell(row=r, column=col).value
                    if val and isinstance(val, str) and kw in val:
                        found = col
                        break
                if found:
                    target_cols.append(found)
                    break
    if not target_cols:
        print("错误：未找到任何学历教师列。")
        return None

    district_col = find_district_col(ws, head_size)   # 新增区县列
    count_col = find_count_school_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '学历教师总数', '每百名学生高学历教师数']
    rows = [header]
    school_index = {}

    for r in range(head_size + 1, max_h + 1):
        stu_val = ws.cell(row=r, column=stu_col).value
        edu_sum = 0
        for c in target_cols:
            v = ws.cell(row=r, column=c).value
            try:
                edu_sum += float(v) if v is not None else 0
            except (ValueError, TypeError):
                pass
        try:
            stu = float(stu_val) if stu_val is not None else None
        except (ValueError, TypeError):
            continue
        if stu is None or stu == 0 or edu_sum == 0:
            continue
        result = round(edu_sum * 100 / stu, 2)
        school_name = ws.cell(row=r, column=1).value
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, edu_sum, result]
        rows.append(row_data)
        school_index[str(school_name).strip()] = len(rows) - 1

    return rows

def calc_backbone_per_stu(ws, head_size, max_h, school_type, merge_map=None):
    """3. 每百名学生骨干教师数 - 不需要调整"""
    if school_type == '小学':
        backbone_col = find_col_by_hierarchy(ws, head_size, ['教基1102', '县级以上骨干教师(小学)'])
    elif school_type == '初中':
        backbone_col = find_col_by_hierarchy(ws, head_size, ['教基1102', '县级以上骨干教师(初中)'])
    else:
        backbone_col = find_col_by_hierarchy(ws, head_size, ['教基1102', '县级以上骨干教师(高中)'])
    if not backbone_col:
        backbone_col = find_col_by_keywords(ws, head_size, '骨干教师')
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    if not stu_col or not backbone_col:
        print("错误：未找到'在校生数'或'骨干教师'列。")
        return None

    district_col = find_district_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '骨干教师数', '每百名学生骨干教师数']
    rows = [header]

    for r in range(head_size + 1, max_h + 1):
        stu_val = ws.cell(row=r, column=stu_col).value
        back_val = ws.cell(row=r, column=backbone_col).value
        try:
            stu = float(stu_val) if stu_val is not None else None
            back = float(back_val) if back_val is not None else None
        except (ValueError, TypeError):
            continue
        if stu is None or back is None or stu == 0:
            continue
        result = round(back * 100 / stu, 2)
        school_name = ws.cell(row=r, column=1).value
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, back, result]
        rows.append(row_data)

    return rows

def calc_art_sport_per_stu(ws, head_size, max_h, merge_map=None):
    """4. 每百名学生艺体教师数 - 不需要调整"""
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    sport_col = find_col_by_hierarchy(ws, head_size, ['专任教师', '#按授课课程分', '体育与健康'])
    art_col = find_col_by_hierarchy(ws, head_size, ['专任教师', '#按授课课程分', '艺术', '计'])
    if not sport_col and not art_col:
        sport_col = find_col_by_keywords(ws, head_size, '体育与健康')
        art_col = find_col_by_keywords(ws, head_size, '艺术')
    if not stu_col or (not sport_col and not art_col):
        print("错误：未找到'在校生数'或体育/艺术相关列。")
        return None

    district_col = find_district_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '体育教师', '艺术教师(计)', '每百名学生艺体教师数']
    rows = [header]

    for r in range(head_size + 1, max_h + 1):
        stu_val = ws.cell(row=r, column=stu_col).value
        sport_val = ws.cell(row=r, column=sport_col).value if sport_col else 0
        art_val = ws.cell(row=r, column=art_col).value if art_col else 0
        try:
            stu = float(stu_val) if stu_val is not None else None
            sport = float(sport_val) if sport_val is not None else 0
            art = float(art_val) if art_val is not None else 0
        except (ValueError, TypeError):
            continue
        if stu is None or stu == 0:
            continue
        total = sport + art
        if total == 0:
            continue
        result = round(total * 100 / stu, 2)
        school_name = ws.cell(row=r, column=1).value
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, sport, art, result]
        rows.append(row_data)

    return rows

# ==================== 需要使用占比的指标（四项） ====================
def calc_media_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map=None):
    """
    1. 每百名学生多媒体教室数
    调整后指标值 = 原始指标值 × 当前学段占比
    """
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    media_col = find_col_by_hierarchy(ws, head_size, ['网络多媒体教室'])
    if not media_col:
        media_col = find_col_by_keywords(ws, head_size, '网络多媒体教室')
    count_col = find_count_school_col(ws, head_size)
    if not stu_col or not media_col:
        print("错误：未找到'在校生数(合计)'或'网络多媒体教室'列。")
        return None

    district_col = find_district_col(ws, head_size)   # 新增区县列

    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '网络多媒体教室', 
              '原始指标值', '学段占比', '调整后指标值', '备注']
    rows = [header]
    school_index = {}
    pending_items = []

    for r in range(head_size + 1, max_h + 1):
        if count_col is not None:
            count_val = ws.cell(row=r, column=count_col).value
            try:
                if float(count_val) == 0:
                    school_name = ws.cell(row=r, column=1).value
                    if not school_name:
                        continue
                    stu_val = ws.cell(row=r, column=stu_col).value
                    media_val = ws.cell(row=r, column=media_col).value
                    district_val = ws.cell(row=r, column=district_col).value if district_col else ''
                    pending_items.append({
                        'row': r,
                        'school': str(school_name).strip(),
                        'district': str(district_val).strip() if district_val else '',
                        'stu': stu_val,
                        'media': media_val
                    })
                    continue
            except (ValueError, TypeError):
                pass

        stu_val = ws.cell(row=r, column=stu_col).value
        media_val = ws.cell(row=r, column=media_col).value
        try:
            stu = float(stu_val) if stu_val is not None else None
            media = float(media_val) if media_val is not None else None
        except (ValueError, TypeError):
            continue
        if stu is None or media is None or stu == 0:
            continue
        
        # 原始指标值
        original_result = round(media * 100 / stu, 2)
        
        # 获取学段占比并计算调整后指标值
        school_name = str(ws.cell(row=r, column=1).value).strip() if ws.cell(row=r, column=1).value else ''
        ratio = get_section_ratio_for_school(school_name, section_ratios, school_type)
        adjusted_result = round(original_result * ratio, 2)
        
        remark = ''
        if school_name in section_ratios:
            ratio_info = section_ratios[school_name]
            remark = f'{ratio_info["type"]}，{school_type}占比{ratio*100:.2f}%'
        
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, media, original_result, round(ratio, 4), adjusted_result, remark]
        rows.append(row_data)
        school_index[school_name] = len(rows) - 1

    for item in pending_items:
        school = item['school']
        if merge_map and school in merge_map:
            master_name = merge_map[school]
        else:
            print(f"\n发现不计校数的学校：{school} (行{item['row']})")
            while True:
                master_name = input(f"请输入该校所属的原校点名称：").strip()
                if not master_name:
                    continue
                if master_name in school_index:
                    break
                print(f"未找到学校'{master_name}'，请重新输入。")

        if master_name in school_index:
            idx = school_index[master_name]
            old_stu = rows[idx][3]   # 注意索引变化：原始行号0, 学校名1, 区县2, 在校生数3, 多媒体4
            old_media = rows[idx][4]
            try:
                add_stu = float(item['stu']) if item['stu'] is not None else 0
                add_media = float(item['media']) if item['media'] is not None else 0
            except (ValueError, TypeError):
                add_stu = add_media = 0
            new_stu = old_stu + add_stu
            new_media = old_media + add_media
            
            original_result = round(new_media * 100 / new_stu, 2) if new_stu > 0 else 0
            ratio = get_section_ratio_for_school(master_name, section_ratios, school_type)
            adjusted_result = round(original_result * ratio, 2)
            
            rows[idx][3] = new_stu
            rows[idx][4] = new_media
            rows[idx][5] = original_result
            rows[idx][6] = round(ratio, 4)
            rows[idx][7] = adjusted_result
            print(f"已将 {school} 合并到 {master_name}")
        else:
            print(f"错误：原校点 {master_name} 未找到，跳过合并 {school}。")
    return rows

def calc_teaching_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map=None):
    """
    5. 生均教学及辅助用房面积
    调整后指标值 = 原始指标值 × 当前学段占比
    """
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    
    teach_row = None
    teach_base_col = None
    for r in range(1, head_size + 1):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(row=r, column=c).value
            if val and isinstance(val, str) and '教学及辅助用房' in val:
                teach_row = r
                teach_base_col = c
                break
        if teach_row:
            break
    
    if not teach_row:
        print("错误：未找到'教学及辅助用房'。")
        return None
    
    teach_area_col = None
    for r in range(teach_row + 1, head_size + 1):
        val = ws.cell(row=r, column=teach_base_col).value
        if val and isinstance(val, str) and val.strip() == '计':
            teach_area_col = teach_base_col
            break
    
    if not teach_area_col:
        for c in range(teach_base_col - 2, teach_base_col + 3):
            if c < 1 or c > ws.max_column:
                continue
            for r in range(teach_row + 1, head_size + 1):
                val = ws.cell(row=r, column=c).value
                if val and isinstance(val, str) and val.strip() == '计':
                    teach_area_col = c
                    break
            if teach_area_col:
                break
    
    if not teach_area_col:
        print("错误：未找到'教学及辅助用房'下方的'计'列。")
        return None
    
    indoor_sport_col = find_col_by_keywords(ws, head_size, '室内体育用房')
    district_col = find_district_col(ws, head_size)
    count_col = find_count_school_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '教学及辅助用房(计)', 
              '室内体育用房', '教学及辅助用房面积(净)', '原始指标值',
              '学段占比', '调整后指标值', '备注']
    rows = [header]
    school_index = {}
    pending_items = []
    
    for r in range(head_size + 1, max_h + 1):
        if count_col is not None:
            count_val = ws.cell(row=r, column=count_col).value
            try:
                if float(count_val) == 0:
                    school_name = ws.cell(row=r, column=1).value
                    if not school_name:
                        continue
                    district_val = ws.cell(row=r, column=district_col).value if district_col else ''
                    pending_items.append({
                        'row': r,
                        'school': str(school_name).strip(),
                        'district': str(district_val).strip() if district_val else '',
                        'stu': ws.cell(row=r, column=stu_col).value,
                        'teach': ws.cell(row=r, column=teach_area_col).value,
                        'indoor': ws.cell(row=r, column=indoor_sport_col).value if indoor_sport_col else 0
                    })
                    continue
            except (ValueError, TypeError):
                pass
        
        school_name = str(ws.cell(row=r, column=1).value).strip() if ws.cell(row=r, column=1).value else ''
        if not school_name:
            continue
            
        stu_val = ws.cell(row=r, column=stu_col).value
        teach_val = ws.cell(row=r, column=teach_area_col).value
        indoor_val = ws.cell(row=r, column=indoor_sport_col).value if indoor_sport_col else 0
        
        try:
            stu = float(stu_val) if stu_val is not None else None
            teach = float(teach_val) if teach_val is not None else 0
            indoor = float(indoor_val) if indoor_val is not None else 0
        except (ValueError, TypeError):
            continue
        
        if stu is None or stu == 0:
            continue
        
        net_area = teach - indoor
        if net_area < 0:
            net_area = 0
        
        original_result = round(net_area / stu, 2)
        ratio = get_section_ratio_for_school(school_name, section_ratios, school_type)
        adjusted_result = round(original_result * ratio, 2)
        
        remark = ''
        if school_name in section_ratios:
            ratio_info = section_ratios[school_name]
            remark = f'{ratio_info["type"]}，{school_type}占比{ratio*100:.2f}%'
        
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, teach, indoor, net_area, original_result, 
                   round(ratio, 4), adjusted_result, remark]
        rows.append(row_data)
        school_index[school_name] = len(rows) - 1
    
    for item in pending_items:
        school = item['school']
        if merge_map and school in merge_map:
            master_name = merge_map[school]
            if master_name in school_index:
                idx = school_index[master_name]
                old_stu = rows[idx][3]   # 索引位置：行号0,学校1,区县2,学生3,教学用房4,室内5,净面积6,原始7...
                old_teach = rows[idx][4]
                old_indoor = rows[idx][5]
                try:
                    add_stu = float(item['stu']) if item['stu'] is not None else 0
                    add_teach = float(item['teach']) if item['teach'] is not None else 0
                    add_indoor = float(item['indoor']) if item['indoor'] is not None else 0
                except (ValueError, TypeError):
                    continue
                new_stu = old_stu + add_stu
                new_teach = old_teach + add_teach
                new_indoor = old_indoor + add_indoor
                net_area = new_teach - new_indoor
                
                original_result = round(net_area / new_stu, 2) if new_stu > 0 else 0
                ratio = get_section_ratio_for_school(master_name, section_ratios, school_type)
                adjusted_result = round(original_result * ratio, 2)
                
                rows[idx][3] = new_stu
                rows[idx][4] = new_teach
                rows[idx][5] = new_indoor
                rows[idx][6] = net_area
                rows[idx][7] = original_result
                rows[idx][8] = round(ratio, 4)
                rows[idx][9] = adjusted_result
                print(f"已将 {school} 合并到 {master_name}")
    
    return rows

def calc_sports_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map=None):
    """
    6. 生均体育运动场馆面积
    调整后指标值 = 原始指标值 × 当前学段占比
    """
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    sports_field_col = find_col_by_keywords(ws, head_size, '运动场地面积')
    if not sports_field_col:
        sports_field_col = find_col_by_keywords(ws, head_size, '运动场地')
    indoor_sport_col = find_col_by_keywords(ws, head_size, '室内体育用房')
    
    if not stu_col or not sports_field_col:
        return None
    
    district_col = find_district_col(ws, head_size)
    count_col = find_count_school_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '运动场地面积', 
              '室内体育用房', '体育运动场馆面积', '原始指标值',
              '学段占比', '调整后指标值', '备注']
    rows = [header]
    school_index = {}
    pending_items = []
    
    for r in range(head_size + 1, max_h + 1):
        if count_col is not None:
            count_val = ws.cell(row=r, column=count_col).value
            try:
                if float(count_val) == 0:
                    school_name = ws.cell(row=r, column=1).value
                    if not school_name:
                        continue
                    district_val = ws.cell(row=r, column=district_col).value if district_col else ''
                    pending_items.append({
                        'row': r,
                        'school': str(school_name).strip(),
                        'district': str(district_val).strip() if district_val else '',
                        'stu': ws.cell(row=r, column=stu_col).value,
                        'sports': ws.cell(row=r, column=sports_field_col).value,
                        'indoor': ws.cell(row=r, column=indoor_sport_col).value if indoor_sport_col else 0
                    })
                    continue
            except (ValueError, TypeError):
                pass
        
        school_name = str(ws.cell(row=r, column=1).value).strip() if ws.cell(row=r, column=1).value else ''
        if not school_name:
            continue
            
        stu_val = ws.cell(row=r, column=stu_col).value
        sports_val = ws.cell(row=r, column=sports_field_col).value
        indoor_val = ws.cell(row=r, column=indoor_sport_col).value if indoor_sport_col else 0
        
        try:
            stu = float(stu_val) if stu_val is not None else None
            sports = float(sports_val) if sports_val is not None else 0
            indoor = float(indoor_val) if indoor_val is not None else 0
        except (ValueError, TypeError):
            continue
        
        if stu is None or stu == 0:
            continue
        
        total_sports_area = sports + indoor
        original_result = round(total_sports_area / stu, 2)
        ratio = get_section_ratio_for_school(school_name, section_ratios, school_type)
        adjusted_result = round(original_result * ratio, 2)
        
        remark = ''
        if school_name in section_ratios:
            ratio_info = section_ratios[school_name]
            remark = f'{ratio_info["type"]}，{school_type}占比{ratio*100:.2f}%'
        
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, sports, indoor, total_sports_area, original_result,
                   round(ratio, 4), adjusted_result, remark]
        rows.append(row_data)
        school_index[school_name] = len(rows) - 1
    
    for item in pending_items:
        school = item['school']
        if merge_map and school in merge_map:
            master_name = merge_map[school]
            if master_name in school_index:
                idx = school_index[master_name]
                old_stu = rows[idx][3]      # 行号0,学校1,区县2,学生3,运动场地4,室内5,总面积6,原始7...
                old_sports = rows[idx][4]
                old_indoor = rows[idx][5]
                try:
                    add_stu = float(item['stu']) if item['stu'] is not None else 0
                    add_sports = float(item['sports']) if item['sports'] is not None else 0
                    add_indoor = float(item['indoor']) if item['indoor'] is not None else 0
                except (ValueError, TypeError):
                    continue
                new_stu = old_stu + add_stu
                new_sports = old_sports + add_sports
                new_indoor = old_indoor + add_indoor
                total_area = new_sports + new_indoor
                
                original_result = round(total_area / new_stu, 2) if new_stu > 0 else 0
                ratio = get_section_ratio_for_school(master_name, section_ratios, school_type)
                adjusted_result = round(original_result * ratio, 2)
                
                rows[idx][3] = new_stu
                rows[idx][4] = new_sports
                rows[idx][5] = new_indoor
                rows[idx][6] = total_area
                rows[idx][7] = original_result
                rows[idx][8] = round(ratio, 4)
                rows[idx][9] = adjusted_result
                print(f"已将 {school} 合并到 {master_name}")
    
    return rows

def calc_equipment_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map=None):
    """
    7. 生均教学仪器设备值
    调整后指标值 = 原始指标值 × 当前学段占比
    """
    stu_col = find_col_by_hierarchy(ws, head_size, ['在校生数', '合计'])
    if not stu_col:
        stu_col = find_col_by_keywords(ws, head_size, '在校生数')
    equipment_col = find_col_by_keywords(ws, head_size, '教学仪器设备资产值(万元)')
    if not equipment_col:
        equipment_col = find_col_by_keywords(ws, head_size, '教学仪器设备')
    
    if not stu_col or not equipment_col:
        return None
    
    district_col = find_district_col(ws, head_size)
    count_col = find_count_school_col(ws, head_size)
    header = ['原始行号', '学校名称', '统计三级', '在校生数(合计)', '教学仪器设备资产值(万元)', 
              '原始指标值', '学段占比', '调整后指标值', '备注']
    rows = [header]
    school_index = {}
    pending_items = []
    
    for r in range(head_size + 1, max_h + 1):
        if count_col is not None:
            count_val = ws.cell(row=r, column=count_col).value
            try:
                if float(count_val) == 0:
                    school_name = ws.cell(row=r, column=1).value
                    if not school_name:
                        continue
                    district_val = ws.cell(row=r, column=district_col).value if district_col else ''
                    pending_items.append({
                        'row': r,
                        'school': str(school_name).strip(),
                        'district': str(district_val).strip() if district_val else '',
                        'stu': ws.cell(row=r, column=stu_col).value,
                        'equipment': ws.cell(row=r, column=equipment_col).value
                    })
                    continue
            except (ValueError, TypeError):
                pass
        
        school_name = str(ws.cell(row=r, column=1).value).strip() if ws.cell(row=r, column=1).value else ''
        if not school_name:
            continue
            
        stu_val = ws.cell(row=r, column=stu_col).value
        equipment_val = ws.cell(row=r, column=equipment_col).value
        
        try:
            stu = float(stu_val) if stu_val is not None else None
            equipment = float(equipment_val) if equipment_val is not None else 0
        except (ValueError, TypeError):
            continue
        
        if stu is None or stu == 0:
            continue
        
        original_result = round(equipment * 10000 / stu, 2)
        ratio = get_section_ratio_for_school(school_name, section_ratios, school_type)
        adjusted_result = round(original_result * ratio, 2)
        
        remark = ''
        if school_name in section_ratios:
            ratio_info = section_ratios[school_name]
            remark = f'{ratio_info["type"]}，{school_type}占比{ratio*100:.2f}%'
        
        district_str = str(ws.cell(row=r, column=district_col).value or '').strip() if district_col else ''
        row_data = [r, school_name, district_str, stu, equipment, original_result, round(ratio, 4), adjusted_result, remark]
        rows.append(row_data)
        school_index[school_name] = len(rows) - 1
    
    for item in pending_items:
        school = item['school']
        if merge_map and school in merge_map:
            master_name = merge_map[school]
            if master_name in school_index:
                idx = school_index[master_name]
                old_stu = rows[idx][3]          # 行号0,学校1,区县2,学生3,设备4,原始5...
                old_equipment = rows[idx][4]
                try:
                    add_stu = float(item['stu']) if item['stu'] is not None else 0
                    add_equipment = float(item['equipment']) if item['equipment'] is not None else 0
                except (ValueError, TypeError):
                    continue
                new_stu = old_stu + add_stu
                new_equipment = old_equipment + add_equipment
                
                original_result = round(new_equipment * 10000 / new_stu, 2) if new_stu > 0 else 0
                ratio = get_section_ratio_for_school(master_name, section_ratios, school_type)
                adjusted_result = round(original_result * ratio, 2)
                
                rows[idx][3] = new_stu
                rows[idx][4] = new_equipment
                rows[idx][5] = original_result
                rows[idx][6] = round(ratio, 4)
                rows[idx][7] = adjusted_result
                print(f"已将 {school} 合并到 {master_name}")
    
    return rows

# ==================== 保存结果 ====================
def save_results_to_file(original_path, school_type, calc_results, section_ratios, all_special_schools):
    """保存分析结果"""
    dir_name = os.path.dirname(original_path)
    base_name = os.path.splitext(os.path.basename(original_path))[0]
    new_name = f"{base_name}_分析结果.xlsx"
    new_path = os.path.join(dir_name, new_name)

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    sheet_names = [
        '每百名学生多媒体教室数',
        '每百名学生高学历教师数',
        '每百名学生骨干教师数',
        '每百名学生艺体教师数',
        '生均教学及辅助用房面积',
        '生均运动场馆面积',
        '生均教学仪器设备值'
    ]
    data_blocks = [
        calc_results.get('media'),
        calc_results.get('high_edu'),
        calc_results.get('backbone'),
        calc_results.get('art_sport'),
        calc_results.get('teaching_area'),
        calc_results.get('sports_area'),
        calc_results.get('equipment')
    ]

    special_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    special_font = Font(color="FF0000", bold=True)
    # 调整后指标值的填充色
    adjusted_fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")

    for sname, block in zip(sheet_names, data_blocks):
        ws_out = out_wb.create_sheet(title=sname)
        if block:
            for row_idx, row in enumerate(block, 1):
                ws_out.append(row)
                # 标记特殊学校
                if len(row) > 1 and row[1] and str(row[1]) in section_ratios:
                    for col in range(1, len(row) + 1):
                        cell = ws_out.cell(row=row_idx, column=col)
                        cell.fill = special_fill
                        cell.font = special_font
        else:
            ws_out.append(['无数据'])
    
    # 添加学段占比分析sheet
    if section_ratios:
        add_section_ratio_sheet(out_wb, section_ratios, all_special_schools)
    
    out_wb.save(new_path)
    print(f"\n分析结果已保存至：{new_path}")
    return new_path

def add_section_ratio_sheet(out_wb, section_ratios, all_special_schools):
    """添加学段占比分析sheet"""
    ws = out_wb.create_sheet(title='学段占比分析')
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    ws.merge_cells('A1:K1')
    ws['A1'].value = '特殊学校学段占比分析（用于四项指标调整）'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.append([])
    ws.append(['说明：以下四项指标使用"调整后指标值 = 原始指标值 × 当前学段占比"进行校正：'])
    ws.append(['  1. 每百名学生多媒体教室数'])
    ws.append(['  2. 生均教学及辅助用房面积'])
    ws.append(['  3. 生均运动场馆面积'])
    ws.append(['  4. 生均教学仪器设备值'])
    ws.append([])
    
    current_row = 8
    
    for special_type in ['九年一贯制', '十二年一贯制', '完全中学']:
        schools = all_special_schools.get(special_type, [])
        if not schools:
            continue
        
        # 权重说明
        weight_desc = {
            '九年一贯制': '权重比例：小学:初中 = 1:1.1',
            '十二年一贯制': '权重比例：小学:初中:高中 = 1:1.1:1.32',
            '完全中学': '权重比例：初中:高中 = 1:1.2'
        }
        
        ws.merge_cells(f'A{current_row}:K{current_row}')
        ws[f'A{current_row}'].value = f'【{special_type}】{weight_desc[special_type]}'
        ws[f'A{current_row}'].font = Font(bold=True, size=12, color="003366")
        ws[f'A{current_row}'].fill = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
        current_row += 1
        
        ratio_sections = []
        if special_type == '九年一贯制':
            ratio_sections = ['小学', '初中']
        elif special_type == '十二年一贯制':
            ratio_sections = ['小学', '初中', '高中']
        else:
            ratio_sections = ['初中', '高中']
        
        headers = ['学校名称', '类型']
        for section in ratio_sections:
            headers.append(f'{section}学生数')
        for section in ratio_sections:
            headers.append(f'{section}权重')
        for section in ratio_sections:
            headers.append(f'{section}加权值')
        headers.append('加权总值')
        for section in ratio_sections:
            headers.append(f'{section}占比(%)')
        
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border
        current_row += 1
        
        for school in schools:
            school_name = school['name']
            ratio_info = section_ratios.get(school_name, {})
            ratios = ratio_info.get('ratios', {})
            section_students = ratio_info.get('section_students', {})
            
            weights_map = {
                '九年一贯制': {'小学': 1, '初中': 1.1},
                '十二年一贯制': {'小学': 1, '初中': 1.1, '高中': 1.32},
                '完全中学': {'初中': 1, '高中': 1.2}
            }
            school_weights = weights_map.get(special_type, {})
            
            row_data = [school_name, special_type]
            for section in ratio_sections:
                row_data.append(section_students.get(section, 0))
            for section in ratio_sections:
                row_data.append(school_weights.get(section, 0))
            for section in ratio_sections:
                row_data.append(round(section_students.get(section, 0) * school_weights.get(section, 0), 2))
            row_data.append(round(ratio_info.get('total_weighted', 0), 2))
            for section in ratio_sections:
                row_data.append(round(ratios.get(section, 0) * 100, 2))
            
            for col_idx, value in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=value)
                cell.border = thin_border
            current_row += 1
        
        current_row += 1
    
    for col in range(1, 13):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 14

# ==================== 文件收集与处理 ====================
def collect_files():
    """收集小学、初中、高中的文件"""
    required_types = ['小学', '初中', '高中']
    file_map = {}
    
    print("=" * 60)
    print("请依次提供小学、初中、高中的 Excel 文件")
    print("=" * 60)
    
    for req_type in required_types:
        while True:
            file_path = input(f"\n请输入{req_type}的 Excel 文件路径：").strip().strip('"').strip("'")
            if file_path.lower() == 'q':
                sys.exit(0)
            if not os.path.exists(file_path):
                print(f"错误：文件不存在 - {file_path}")
                continue
            file_map[req_type] = file_path
            print(f"已添加{req_type}文件：{os.path.basename(file_path)}")
            break
    
    return file_map

def process_single_file(file_path, school_type, section_ratios):
    """处理单个文件"""
    wb, ws, max_h, max_l = get_worksheet_info(file_path)
    if wb is None:
        return None
    
    print(f"\n{'='*60}")
    print(f"处理文件：{os.path.basename(file_path)}  类型：{school_type}")
    
    head_size, _ = guess_head_size(ws)
    print(f"自动检测表头行数：{head_size}")
    
    user_input = input("确认表头行数（回车确认 / 输入正确行数）：").strip()
    if user_input.isdigit():
        head_size = int(user_input)
    
    merge_map = build_school_merge_map(ws, head_size)
    
    calc_results = {
        'media': None, 'high_edu': None, 'backbone': None,
        'art_sport': None, 'teaching_area': None, 'sports_area': None,
        'equipment': None,
    }
    
    while True:
        print("\n请选择计算项目：")
        print("  1. 每百名学生多媒体教室数（需要学段占比调整）")
        print("  2. 每百名学生高学历教师数")
        print("  3. 每百名学生骨干教师数")
        print("  4. 每百名学生艺体教师数")
        print("  5. 生均教学及辅助用房面积（需要学段占比调整）")
        print("  6. 生均运动场馆面积（需要学段占比调整）")
        print("  7. 生均教学仪器设备值（需要学段占比调整）")
        print("  8. 一键执行以上全部")
        print("  0. 完成此文件，生成报告")
        choice = input("请输入选项：").strip()
        
        if choice == '0':
            break
        elif choice == '8':
            calc_results['media'] = calc_media_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
            calc_results['high_edu'] = calc_high_education_per_stu(ws, head_size, max_h, school_type, merge_map)
            calc_results['backbone'] = calc_backbone_per_stu(ws, head_size, max_h, school_type, merge_map)
            calc_results['art_sport'] = calc_art_sport_per_stu(ws, head_size, max_h, merge_map)
            calc_results['teaching_area'] = calc_teaching_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
            calc_results['sports_area'] = calc_sports_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
            calc_results['equipment'] = calc_equipment_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
            print("全部计算完成。")
            break
        elif choice == '1':
            calc_results['media'] = calc_media_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
        elif choice == '2':
            calc_results['high_edu'] = calc_high_education_per_stu(ws, head_size, max_h, school_type, merge_map)
        elif choice == '3':
            calc_results['backbone'] = calc_backbone_per_stu(ws, head_size, max_h, school_type, merge_map)
        elif choice == '4':
            calc_results['art_sport'] = calc_art_sport_per_stu(ws, head_size, max_h, merge_map)
        elif choice == '5':
            calc_results['teaching_area'] = calc_teaching_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
        elif choice == '6':
            calc_results['sports_area'] = calc_sports_area_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
        elif choice == '7':
            calc_results['equipment'] = calc_equipment_per_stu_with_ratio(ws, head_size, max_h, school_type, section_ratios, merge_map)
        else:
            print("无效输入。")
    
    return calc_results

def main():
    print("=" * 60)
    print("Excel 教育数据综合分析工具")
    print("支持特殊学校学段占比计算与指标调整")
    print("=" * 60)
    print("\n指标调整说明：")
    print("  四项指标需要调整：多媒体教室数、教学用房面积、运动场馆面积、教学仪器设备值")
    print("  三项指标无需调整：高学历教师数、骨干教师数、艺体教师数")
    print("  调整方式：调整后指标值 = 原始指标值 × 当前学段占比")
    
    file_map = collect_files()
    
    # ===== 第一步：查找所有特殊学校 =====
    print("\n" + "=" * 60)
    print("第一步：查找特殊学校类型")
    print("=" * 60)
    
    all_special_schools = {
        '九年一贯制': [],
        '十二年一贯制': [],
        '完全中学': []
    }
    
    for school_type in ['小学', '初中', '高中']:
        file_path = file_map[school_type]
        wb, ws, max_h, _ = get_worksheet_info(file_path)
        if wb is None:
            continue
        head_size, _ = guess_head_size(ws)
        special_schools = find_special_school_types(ws, head_size, max_h, school_type)
        if special_schools:
            for stype in all_special_schools:
                if stype in special_schools:
                    existing_names = {s['name'] for s in all_special_schools[stype]}
                    for school in special_schools[stype]:
                        if school['name'] not in existing_names:
                            all_special_schools[stype].append(school)
    
    # 去重
    for stype in all_special_schools:
        seen = set()
        unique = []
        for school in all_special_schools[stype]:
            if school['name'] not in seen:
                seen.add(school['name'])
                unique.append(school)
        all_special_schools[stype] = unique
    
    total = sum(len(schools) for schools in all_special_schools.values())
    print(f"\n共找到 {total} 所特殊学校")
    
    # ===== 第二步：计算学段占比 =====
    section_ratios = calculate_section_ratios(all_special_schools, file_map)
    
    # ===== 第三步：处理每个文件 =====
    print("\n" + "=" * 60)
    print("第三步：处理各文件（四项指标使用学段占比调整）")
    print("=" * 60)
    
    for school_type in ['小学', '初中', '高中']:
        file_path = file_map[school_type]
        calc_results = process_single_file(file_path, school_type, section_ratios)

        if calc_results and any(v is not None for v in calc_results.values()):
            new_path = save_results_to_file(file_path, school_type, calc_results, section_ratios, all_special_schools)

            # 询问是否进行区县差异系数统计
            district_choice = input(f"\n是否对文件 {os.path.basename(file_path)} 进行区县差异系数统计？(y/n)：").strip().lower()
            if district_choice == 'y':
                stats_rows, summary_rows = analyze_district_from_result_file(new_path, school_type)
                if stats_rows or summary_rows:
                    append_district_sheets_to_file(new_path, stats_rows, summary_rows)

    # 循环结束后才打印处理完成
    print("\n" + "=" * 60)
    print("处理完成！")
    print("=" * 60)
    
    # 显示调整汇总
    if section_ratios:
        print("\n特殊学校指标调整汇总：")
        for school_name, info in section_ratios.items():
            ratios = info['ratios']
            print(f"  {school_name} ({info['type']})：")
            for section, ratio in ratios.items():
                print(f"    {section}占比：{ratio*100:.2f}%")
    
    input("\n按 Enter 键退出...")

if __name__ == "__main__":
    main()